#!/usr/bin/env python3
"""Build an Excel workbook with two dashboard sheets and the detail tables."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.chart.data_source import AxDataSource, StrRef
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import Marker
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.comments import Comment
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import (
    CharacterProperties,
    Font as DrawingFont,
    Paragraph,
    ParagraphProperties,
    RichTextProperties,
)
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart.legend import Legend

from build_data import VALID_STATUS, load_orders, load_week, WEEKS

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(__file__).resolve().parent / "营多天猫经营仪表盘.xlsx"

GROUP_ORDER = ["袋面主链接", "整箱/大包", "杯面", "袋面次链接", "口味与组合"]
GROUP_COLORS = ["E879F9", "A78BFA", "818CF8", "F0ABFC", "67E8F9"]

PURPLE_BG = "100818"
PURPLE_CARD = "24143F"
TEAL_BG = "071816"
TEAL_CARD = "0E2E2C"
WHITE = "F7F4FF"
MUTED = "C4B6E4"
TEAL_MUTED = "A7CDC6"
PINK = "FB7185"
TEAL = "5EEAD4"
INK = "1E1633"
LINE = "3A3058"
HEADER = "3B0764"
ZEBRA = "F6F3FB"
INPUT_BG = "F8F5FF"
INPUT_FG = "2E1064"


def font(size=11, bold=False, color=WHITE, name="微软雅黑"):
    return Font(name=name, size=size, bold=bold, color=color)


def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)


def align(horizontal="left", vertical="center", wrap=False):
    return Alignment(horizontal=horizontal, vertical=vertical, wrap_text=wrap)


def side(color):
    return Side(style="thin", color=color)


def apply_range(ws, row1, col1, row2, col2, *, fg=None, color=None, size=11, bold=False, horizontal="left", wrap=False, border_color=None):
    for row in range(row1, row2 + 1):
        for col in range(col1, col2 + 1):
            cell = ws.cell(row, col)
            if fg:
                cell.fill = fill(fg)
            cell.font = font(size, bold, color or WHITE)
            cell.alignment = align(horizontal, wrap=wrap)
            if border_color:
                cell.border = Border(
                    left=side(border_color) if col == col1 else side(fg or border_color),
                    right=side(border_color) if col == col2 else side(fg or border_color),
                    top=side(border_color) if row == row1 else side(fg or border_color),
                    bottom=side(border_color) if row == row2 else side(fg or border_color),
                )


def paint_sheet(ws, rows, cols, fg):
    for row in range(1, rows + 1):
        for col in range(1, cols + 1):
            cell = ws.cell(row, col)
            cell.fill = fill(fg)
            cell.font = font(10, color=MUTED)


def text_props(color="F4F1FF", size=900):
    face = DrawingFont(typeface="微软雅黑")
    char = CharacterProperties(solidFill=color, sz=size, latin=face, ea=face)
    return RichText(
        bodyPr=RichTextProperties(anchor="ctr"),
        p=[Paragraph(pPr=ParagraphProperties(defRPr=char), endParaRPr=char, r=[])],
    )


def tint_title(chart, color="F7F4FF"):
    title = chart.title
    if title is None or title.tx is None or title.tx.rich is None:
        return
    face = DrawingFont(typeface="微软雅黑")
    char = CharacterProperties(solidFill=color, sz=1200, b=True, latin=face, ea=face)
    for paragraph in title.tx.rich.p:
        if paragraph.pPr is None:
            paragraph.pPr = ParagraphProperties()
        paragraph.pPr.defRPr = char
        paragraph.endParaRPr = char
        for run in paragraph.r:
            run.rPr = char


def style_axis(axis, color="F4F1FF"):
    axis.txPr = text_props(color, 900)
    axis.spPr = GraphicalProperties(ln=LineProperties(solidFill="6B5B8A", w=6000))
    if axis.majorGridlines is not None:
        axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="3A3058", w=4000))


def style_dark_chart(chart, plot="1B1430"):
    chart.graphical_properties = GraphicalProperties(solidFill="120C22")
    if chart.plot_area.graphicalProperties is None:
        chart.plot_area.graphicalProperties = GraphicalProperties(solidFill=plot)
    else:
        chart.plot_area.graphicalProperties.solidFill = plot
    if chart.legend is None:
        chart.legend = Legend()
    chart.legend.position = "b"
    chart.legend.overlay = False
    chart.legend.txPr = text_props("F4F1FF", 900)
    chart.legend.graphicalProperties = GraphicalProperties(solidFill="120C22")
    chart.visible_cells_only = False
    tint_title(chart)


def use_text_categories(chart, cats):
    chart.set_categories(cats)
    formula = str(cats)
    for series in chart.series:
        series.cat = AxDataSource(strRef=StrRef(formula))


def groups_for(week):
    merged = {}
    for group in week["groups"]:
        name = "口味与组合" if (group["share"] or 0) < 1 else group["name"]
        merged[name] = merged.get(name, 0.0) + group["gmv"]
    total = week["totals"]["gmv"] or 1
    rows = []
    for name in GROUP_ORDER:
        gmv = round(merged.get(name, 0.0), 2)
        rows.append((name, gmv, round(gmv / total * 100, 2)))
    return rows


def short_prov(name):
    return str(name).replace("省", "").replace("市", "").replace("壮族自治区", "").replace("回族自治区", "").replace("维吾尔自治区", "").replace("自治区", "")


def week_rows(weeks):
    rows = []
    for week in weeks:
        totals = week["totals"]
        groups = {name: gmv for name, gmv, _share in groups_for(week)}
        rows.append(
            {
                "周次": week["label"],
                "开始": week["start"],
                "结束": week["end"],
                "备注": week["note"] or "",
                "成交金额": totals["gmv"],
                "去税金额": totals["net"],
                "成交件数": totals["qty"],
                "件单价": totals["asp"],
                "推广总花费": totals["adSpend"],
                "推广总成交": totals["adSales"],
                "推广ROI": totals["adRoi"],
                "万相台花费": totals["wxtSpend"],
                "万相台成交": totals["wxtSales"],
                "万相台ROI": totals["wxtRoi"],
                "品销宝花费": totals["pxbSpend"],
                "品销宝成交": totals["pxbSales"],
                "品销宝ROI": totals["pxbRoi"],
                "日销成交": totals["organic"],
                "其他费用": totals["other"],
                "推广前利润": totals["profitBefore"],
                "推广后利润": totals["profitAfter"],
                "推广前净利率": totals["marginBefore"],
                "推广后净利率": totals["marginAfter"],
                "推广费率": totals["adRate"],
                "日销占比": totals["organicShare"],
                "万相台覆盖": totals["wxtCover"],
                "推广覆盖": totals["adCover"],
                "主链接占比": totals["mainShare"],
                "前三集中度": totals["top3Share"],
                "杯面主链接占比": totals["cupShare"],
                "链接数": totals["skuCount"],
                **{name: groups[name] for name in GROUP_ORDER},
            }
        )
    return rows


WEEK_FIELDS = [
    "周次", "开始", "结束", "备注",
    "成交金额", "去税金额", "成交件数", "件单价",
    "推广总花费", "推广总成交", "推广ROI",
    "万相台花费", "万相台成交", "万相台ROI",
    "品销宝花费", "品销宝成交", "品销宝ROI",
    "日销成交", "其他费用",
    "推广前利润", "推广后利润",
    "推广前净利率", "推广后净利率", "推广费率", "日销占比",
    "万相台覆盖", "推广覆盖", "主链接占比", "前三集中度", "杯面主链接占比", "链接数",
    *GROUP_ORDER,
]


def write_week_sheet(wb, rows):
    ws = wb.create_sheet("周汇总")
    ws.sheet_properties.tabColor = "6D28D9"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(WEEK_FIELDS))}{len(rows) + 1}"
    ws.row_dimensions[1].height = 22
    percent_fields = {"推广前净利率", "推广后净利率", "推广费率", "日销占比", "万相台覆盖", "推广覆盖", "主链接占比", "前三集中度", "杯面主链接占比"}
    roi_fields = {"推广ROI", "万相台ROI", "品销宝ROI"}
    int_fields = {"成交件数", "链接数"}
    for col, name in enumerate(WEEK_FIELDS, 1):
        cell = ws.cell(1, col, name)
        cell.fill = fill(HEADER)
        cell.font = font(10, True, WHITE)
        cell.alignment = align("center")
        ws.column_dimensions[get_column_letter(col)].width = 14 if name != "备注" else 42
    ws.column_dimensions["A"].width = 14
    for r, row in enumerate(rows, 2):
        ws.row_dimensions[r].height = 18
        for c, name in enumerate(WEEK_FIELDS, 1):
            cell = ws.cell(r, c, row[name] if row[name] != "" else None)
            cell.font = font(10, color=INK)
            cell.fill = fill(ZEBRA if r % 2 == 0 else "FFFFFF")
            cell.alignment = align("center" if c <= 3 else "right")
            if name in percent_fields:
                cell.number_format = '0.0"%"'
            elif name in roi_fields:
                cell.number_format = "0.00"
            elif name in int_fields:
                cell.number_format = "#,##0"
            elif isinstance(row[name], float):
                cell.number_format = '#,##0.00'
    ws.auto_filter.ref = f"A1:{get_column_letter(len(WEEK_FIELDS))}{len(rows) + 1}"
    ws.oddHeader.left.text = "营多天猫 · 周汇总"
    return ws


def col_letter(name):
    return get_column_letter(WEEK_FIELDS.index(name) + 1)


def week_index_formula(column_name, week_cell="$C$3", last_row=9):
    letter = col_letter(column_name)
    return f'INDEX(周汇总!{letter}$2:{letter}${last_row},MATCH({week_cell},周汇总!$A$2:$A${last_row},0))'


def write_chart_data(wb, weeks, orders, week_last):
    ws = wb.create_sheet("图表数据")
    ws.sheet_state = "hidden"
    headers = ["店铺", "订单数", "件数", "商家实收", "买家实付", "去税收入", "总成本", "推广前利润", "净利率"]
    for col, name in enumerate(headers, 1):
        ws.cell(1, col, name)
    for row_index, shop in enumerate(orders["shops"], 2):
        values = [shop["name"], shop["orders"], shop["qty"], shop["recv"], shop["pay"], shop["net"], shop["cost"], shop["profit"], shop["margin"]]
        for col, value in enumerate(values, 1):
            ws.cell(row_index, col, value)
    ws.cell(4, 1, "合计")
    for col in range(2, 9):
        letter = get_column_letter(col)
        ws.cell(4, col, f"={letter}2+{letter}3")
    ws.cell(4, 9, "=IF(F4=0,NA(),H4/F4*100)")

    daily_header_row = 6
    daily_headers = ["日期", "旗舰实收", "淘宝实收", "旗舰件数", "淘宝件数", "旗舰成本", "淘宝成本", "旗舰利润", "淘宝利润", "日期标签", "图_实收1", "图_实收2"]
    for col, name in enumerate(daily_headers, 1):
        ws.cell(daily_header_row, col, name)
    for offset, day in enumerate(orders["daily"]):
        row = daily_header_row + 1 + offset
        month, day_num = int(day["date"][5:7]), int(day["date"][8:10])
        values = [
            day["date"], day["flagship"], day["taobao"], day["qtyFlagship"], day["qtyTaobao"],
            day["costFlagship"], day["costTaobao"], day["profitFlagship"], day["profitTaobao"],
            f"{month}/{day_num}",
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row, col, value)
        ws.cell(row, 11, f'=IF(双店销售!$C$3="淘宝店",C{row},B{row})')
        ws.cell(row, 12, f'=IF(双店销售!$C$3="两家店",C{row},NA())')
    daily_first = daily_header_row + 1
    daily_last = daily_header_row + len(orders["daily"])
    ws.cell(daily_header_row, 11, '=IF(双店销售!$C$3="淘宝店","淘宝店","旗舰店")')
    ws.cell(daily_header_row, 12, '=IF(双店销售!$C$3="两家店","淘宝店","不显示")')

    prov_row = 16
    ws.cell(prov_row, 1, "省份")
    ws.cell(prov_row, 2, "旗舰店")
    ws.cell(prov_row, 3, "淘宝店")
    ws.cell(prov_row, 4, "图_金额1")
    ws.cell(prov_row, 5, "图_金额2")
    for offset, item in enumerate(orders["provinces"]):
        row = prov_row + 1 + offset
        ws.cell(row, 1, short_prov(item["name"]))
        ws.cell(row, 2, item["flagship"])
        ws.cell(row, 3, item["taobao"])
        ws.cell(row, 4, f'=IF(双店销售!$C$3="淘宝店",C{row},B{row})')
        ws.cell(row, 5, f'=IF(双店销售!$C$3="两家店",C{row},NA())')
    prov_first = prov_row + 1
    prov_last = prov_row + len(orders["provinces"])
    ws.cell(prov_row, 4, '=IF(双店销售!$C$3="淘宝店","淘宝店","旗舰店")')
    ws.cell(prov_row, 5, '=IF(双店销售!$C$3="两家店","淘宝店","不显示")')

    series_row = 28
    ws.cell(series_row, 1, "系列")
    ws.cell(series_row, 2, "旗舰店")
    ws.cell(series_row, 3, "淘宝店")
    ws.cell(series_row, 4, "图_金额1")
    ws.cell(series_row, 5, "图_金额2")
    for offset, item in enumerate(orders["series"]):
        row = series_row + 1 + offset
        ws.cell(row, 1, item["name"])
        ws.cell(row, 2, item["flagship"])
        ws.cell(row, 3, item["taobao"])
        ws.cell(row, 4, f'=IF(双店销售!$C$3="淘宝店",C{row},B{row})')
        ws.cell(row, 5, f'=IF(双店销售!$C$3="两家店",C{row},NA())')
    series_first = series_row + 1
    series_last = series_row + len(orders["series"])
    ws.cell(series_row, 4, '=IF(双店销售!$C$3="淘宝店","淘宝店","旗舰店")')
    ws.cell(series_row, 5, '=IF(双店销售!$C$3="两家店","淘宝店","不显示")')

    matrix_row = 38
    ws.cell(matrix_row, 1, "周次")
    for col, name in enumerate(GROUP_ORDER, 2):
        ws.cell(matrix_row, col, name)
    for offset, week in enumerate(weeks):
        row = matrix_row + 1 + offset
        ws.cell(row, 1, week["label"])
        grouped = {name: gmv for name, gmv, _share in groups_for(week)}
        for col, name in enumerate(GROUP_ORDER, 2):
            ws.cell(row, col, grouped[name])
    matrix_first = matrix_row + 1
    matrix_last = matrix_row + len(weeks)

    pick_row = 48
    ws.cell(pick_row, 1, "结构")
    ws.cell(pick_row, 2, "成交金额")
    for offset, name in enumerate(GROUP_ORDER):
        row = pick_row + 1 + offset
        letter = get_column_letter(offset + 2)
        ws.cell(row, 1, name)
        ws.cell(row, 2, f'=INDEX({letter}${matrix_first}:{letter}${matrix_last},MATCH(经营总览!$C$3,$A${matrix_first}:$A${matrix_last},0))')
    pick_first = pick_row + 1
    pick_last = pick_row + len(GROUP_ORDER)

    return {
        "daily_header": daily_header_row,
        "daily_first": daily_first,
        "daily_last": daily_last,
        "prov_header": prov_row,
        "prov_first": prov_first,
        "prov_last": prov_last,
        "series_header": series_row,
        "series_first": series_first,
        "series_last": series_last,
        "pick_header": pick_row,
        "pick_first": pick_first,
        "pick_last": pick_last,
        "week_last": week_last,
    }


def add_combo(ws, anchor, title, cats, bar_data, line_data):
    bar = BarChart()
    bar.type = "col"
    bar.grouping = "clustered"
    bar.title = title
    bar.y_axis.axId = 100
    bar.y_axis.crosses = "min"
    bar.y_axis.scaling.min = 0
    bar.y_axis.numFmt = "#,##0"
    bar.gapWidth = 60
    bar.add_data(bar_data, titles_from_data=True)
    use_text_categories(bar, cats)
    bar.x_axis.axPos = "b"
    bar.y_axis.axPos = "l"
    bar.legend.position = "b"

    line = LineChart()
    line.y_axis.axId = 200
    line.y_axis.axPos = "r"
    line.y_axis.crosses = "max"
    line.y_axis.numFmt = "#,##0"
    line.y_axis.majorGridlines = None
    line.x_axis.axId = bar.x_axis.axId
    line.smooth = True
    line.add_data(line_data, titles_from_data=True)
    use_text_categories(line, cats)

    bar.series[0].graphicalProperties.solidFill = "A855F7"
    line.series[0].graphicalProperties.line.solidFill = "67E8F9"
    line.series[0].graphicalProperties.line.width = 25000
    line.series[0].marker = Marker(symbol="circle", size=7)
    line.series[0].marker.graphicalProperties.solidFill = "E0F2FE"
    line.series[0].marker.graphicalProperties.line.solidFill = "67E8F9"
    line.series[0].smooth = True

    bar += line
    style_axis(bar.x_axis)
    style_axis(bar.y_axis)
    style_axis(line.y_axis)
    style_dark_chart(bar)
    bar.width = 15.2
    bar.height = 7.6
    ws.add_chart(bar, anchor)


def add_doughnut(ws, anchor, cats, data):
    chart = DoughnutChart()
    chart.title = "成交结构"
    chart.add_data(data, titles_from_data=True)
    use_text_categories(chart, cats)
    chart.holeSize = 62
    chart.series[0].data_points = [
        DataPoint(idx=index, spPr=GraphicalProperties(solidFill=color))
        for index, color in enumerate(GROUP_COLORS)
    ]
    labels = DataLabelList()
    labels.showPercent = True
    labels.showVal = False
    labels.showCatName = False
    labels.showSerName = False
    labels.numFmt = "0.0%"
    labels.txPr = text_props("F7F4FF", 900)
    chart.series[0].dLbls = labels
    style_dark_chart(chart)
    chart.legend.position = "r"
    chart.width = 13.6
    chart.height = 8.0
    ws.add_chart(chart, anchor)


def add_clustered(ws, anchor, title, cats, data, colors, horizontal=False, width=15, height=7.2, largest_first=False):
    chart = BarChart()
    chart.type = "bar" if horizontal else "col"
    chart.grouping = "clustered"
    chart.title = title
    chart.y_axis.numFmt = "#,##0"
    chart.gapWidth = 70
    chart.overlap = -8
    chart.add_data(data, titles_from_data=True)
    use_text_categories(chart, cats)
    chart.x_axis.axPos = "l" if horizontal else "b"
    chart.y_axis.axPos = "b" if horizontal else "l"
    if largest_first:
        chart.x_axis.scaling.orientation = "maxMin"
    for series, color in zip(chart.series, colors):
        series.graphicalProperties.solidFill = color
    style_axis(chart.x_axis)
    style_axis(chart.y_axis)
    style_dark_chart(chart, "102624" if not horizontal else "102624")
    chart.width = width
    chart.height = height
    ws.add_chart(chart, anchor)


def add_lines(ws, anchor, title, cats, data):
    chart = LineChart()
    chart.title = title
    chart.y_axis.scaling.min = 0
    chart.y_axis.numFmt = "#,##0"
    chart.add_data(data, titles_from_data=True)
    use_text_categories(chart, cats)
    chart.x_axis.axPos = "b"
    chart.y_axis.axPos = "l"
    colors = ["5EEAD4", "F9A8D4"]
    for series, color in zip(chart.series, colors):
        series.graphicalProperties.line.solidFill = color
        series.graphicalProperties.line.width = 25000
        series.marker = Marker(symbol="circle", size=6)
        series.marker.graphicalProperties.solidFill = color
        series.marker.graphicalProperties.line.solidFill = color
        series.smooth = True
    style_axis(chart.x_axis, "E7FFFA")
    style_axis(chart.y_axis, "E7FFFA")
    style_dark_chart(chart, "0C2422")
    chart.graphical_properties = GraphicalProperties(solidFill="071816")
    chart.width = 15.2
    chart.height = 7.4
    ws.add_chart(chart, anchor)


def setup_dashboard(ws, bg, tab):
    ws.sheet_properties.tabColor = tab
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 110
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.horizontalCentered = True
    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.4
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.4
    ws.sheet_view.view = "normal"
    ws.print_options.horizontalCentered = True
    paint_sheet(ws, 46, 16, bg)
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["N"].width = 3
    ws.row_dimensions[1].height = 8
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[3].height = 22
    ws.row_dimensions[4].height = 8
    for row in range(5, 11):
        ws.row_dimensions[row].height = 20
    ws.row_dimensions[6].height = 26
    ws.row_dimensions[7].height = 26
    widths = {
        "B": 18, "C": 14, "D": 10,
        "E": 18, "F": 14, "G": 10,
        "H": 18, "I": 14, "J": 10,
        "K": 18, "L": 14, "M": 4,
    }
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width


def write_card(ws, col, title, hero_formula, hero_format, lines, card_fg, accent, muted=MUTED, hero_color=WHITE):
    """lines: list of (label, formula, number_format). Card uses two columns; the next column stays as a gutter."""
    apply_range(ws, 5, col, 10, col + 1, fg=card_fg, border_color=accent, color=hero_color)
    ws.merge_cells(start_row=5, start_column=col, end_row=5, end_column=col + 1)
    ws.merge_cells(start_row=6, start_column=col, end_row=7, end_column=col + 1)
    title_cell = ws.cell(5, col, title)
    title_cell.font = font(10, True, muted)
    title_cell.alignment = align("left")
    hero = ws.cell(6, col, hero_formula)
    hero.font = font(22, True, hero_color)
    hero.alignment = align("left")
    hero.number_format = hero_format
    for offset, (label, formula, number_format) in enumerate(lines):
        row = 8 + offset
        label_cell = ws.cell(row, col, label)
        label_cell.font = font(9, False, muted)
        label_cell.fill = fill(card_fg)
        label_cell.alignment = align("left")
        value = ws.cell(row, col + 1, formula)
        value.font = font(11, True, hero_color)
        value.alignment = align("right")
        value.number_format = number_format
        value.fill = fill(card_fg)


def write_overview(wb, refs):
    ws = wb.create_sheet("经营总览", 0)
    setup_dashboard(ws, PURPLE_BG, "C026D3")
    last = refs["week_last"]
    ws.merge_cells("B2:M2")
    title = ws.cell(2, 2, "营多捞面 · 天猫经营总览")
    title.font = font(22, True, WHITE)
    title.alignment = align("left")
    ws.cell(3, 2, "周次").font = font(11, True, MUTED)
    ws.merge_cells("C3:D3")
    picker = ws.cell(3, 3, wb["周汇总"]["A9"].value)
    picker.font = font(12, True, INPUT_FG)
    picker.fill = fill(INPUT_BG)
    picker.alignment = align("center")
    ws.cell(3, 4).fill = fill(INPUT_BG)
    picker.comment = Comment("选择周次后，四张指标卡和右侧结构图会改成该周。下方柱线趋势图固定显示全部 8 个周。", "仪表盘")
    ws.merge_cells("E3:M3")
    note = ws.cell(3, 5, f'=IF({week_index_formula("备注", "$C$3", last)}="", "指标随周次变化。趋势图为全部周次。", {week_index_formula("备注", "$C$3", last)})')
    note.font = font(9, False, MUTED)
    validation = DataValidation(type="list", formula1="=周次列表", allow_blank=False)
    validation.error = "请从列表中选择周次"
    validation.errorTitle = "周次"
    validation.prompt = "选择一个周次"
    validation.promptTitle = "周次"
    validation.add(ws["C3"])
    ws.add_data_validation(validation)

    money = '#,##0'
    pct = '0.0"%"'
    roi = "0.00"
    write_card(
        ws, 2, "成交",
        f'={week_index_formula("成交金额", "$C$3", last)}', money,
        [
            ("去税金额", f'={week_index_formula("去税金额", "$C$3", last)}', money),
            ("成交件数", f'={week_index_formula("成交件数", "$C$3", last)}', "#,##0"),
            ("件单价", f'={week_index_formula("件单价", "$C$3", last)}', "0.00"),
        ],
        PURPLE_CARD, "6D28D9",
    )
    write_card(
        ws, 5, "盈利能力",
        f'={week_index_formula("推广前净利率", "$C$3", last)}', pct,
        [
            ("推广后净利率", f'={week_index_formula("推广后净利率", "$C$3", last)}', '0.0"%";-0.0"%"'),
            ("推广前利润", f'={week_index_formula("推广前利润", "$C$3", last)}', money),
            ("较上周(pt)", f'=IF(MATCH($C$3,周汇总!$A$2:$A${last},0)=1,"—",{week_index_formula("推广前净利率", "$C$3", last)}-INDEX(周汇总!{col_letter("推广前净利率")}$2:{col_letter("推广前净利率")}${last},MATCH($C$3,周汇总!$A$2:$A${last},0)-1))', '0.00;-0.00'),
        ],
        PURPLE_CARD, "C026D3",
    )
    write_card(
        ws, 8, "推广效率",
        f'={week_index_formula("推广ROI", "$C$3", last)}', roi,
        [
            ("万相台 ROI", f'={week_index_formula("万相台ROI", "$C$3", last)}', roi),
            ("推广费率", f'={week_index_formula("推广费率", "$C$3", last)}', pct),
            ("万相台覆盖", f'={week_index_formula("万相台覆盖", "$C$3", last)}', pct),
        ],
        PURPLE_CARD, "7C3AED",
    )
    write_card(
        ws, 11, "链接结构",
        f'={week_index_formula("主链接占比", "$C$3", last)}', pct,
        [
            ("前三集中度", f'={week_index_formula("前三集中度", "$C$3", last)}', pct),
            ("日销占比", f'={week_index_formula("日销占比", "$C$3", last)}', '0.0"%";-0.0"%"'),
            ("在售链接", f'={week_index_formula("链接数", "$C$3", last)}', "0"),
        ],
        "5B21B6", "E9D5FF",
    )

    ws.merge_cells("B12:G12")
    left_title = ws.cell(12, 2, "成交金额（柱，左轴）与推广后净利润（线，右轴）")
    left_title.font = font(11, True, WHITE)
    ws.merge_cells("H12:M12")
    right_title = ws.cell(12, 8, "所选周次的成交结构")
    right_title.font = font(11, True, WHITE)

    data = wb["周汇总"]
    cats = Reference(data, min_col=1, min_row=2, max_row=last)
    gmv_col = WEEK_FIELDS.index("成交金额") + 1
    profit_col = WEEK_FIELDS.index("推广后利润") + 1
    add_combo(
        ws, "B13", "八周趋势",
        cats,
        Reference(data, min_col=gmv_col, min_row=1, max_row=last),
        Reference(data, min_col=profit_col, min_row=1, max_row=last),
    )
    source = wb["图表数据"]
    add_doughnut(
        ws, "H13",
        Reference(source, min_col=1, min_row=refs["pick_first"], max_row=refs["pick_last"]),
        Reference(source, min_col=2, min_row=refs["pick_header"], max_row=refs["pick_last"]),
    )

    apply_range(ws, 31, 2, 37, 4, fg=PURPLE_CARD, border_color="6D28D9", color=WHITE)
    structure_title = ws.cell(31, 2, "结构金额")
    structure_title.font = font(11, True, WHITE)
    ws.merge_cells("B31:D31")
    headers = ["名称", "金额", "占比"]
    for col, name in enumerate(headers, 2):
        cell = ws.cell(32, col, name)
        cell.font = font(9, True, MUTED)
        cell.fill = fill(PURPLE_CARD)
    for offset in range(5):
        row = 33 + offset
        src = refs["pick_first"] + offset
        name_cell = ws.cell(row, 2, f"=图表数据!A{src}")
        name_cell.font = font(10, False, WHITE)
        name_cell.fill = fill(PURPLE_CARD)
        amount = ws.cell(row, 3, f"=图表数据!B{src}")
        amount.font = font(10, True, WHITE)
        amount.number_format = '#,##0'
        amount.fill = fill(PURPLE_CARD)
        share = ws.cell(row, 4, f"=IF(SUM(图表数据!$B${refs['pick_first']}:$B${refs['pick_last']})=0,NA(),C{row}/SUM(图表数据!$B${refs['pick_first']}:$B${refs['pick_last']}))")
        share.font = font(10, True, WHITE)
        share.number_format = "0.0%"
        share.fill = fill(PURPLE_CARD)

    ws.merge_cells("B38:M40")
    footnote = ws.cell(
        38, 2,
        "净利率 = 净利润 ÷ 去税金额。推广费率、日销占比、成交覆盖 = 对应金额 ÷ 成交金额。ROI = 成交 ÷ 花费。"
        "八个周的推广后利润都是负的，推广前净利率约 15%。缺周不补零。"
        "7.6–7.12 的源文件日期列误写，周次按文件名。7.27–8.2 的源文件日期列写为 7.20–8.2。",
    )
    footnote.font = font(9, False, MUTED)
    footnote.alignment = align("left", "center", True)
    ws.row_dimensions[38].height = 18
    ws.row_dimensions[39].height = 18
    ws.row_dimensions[40].height = 18
    ws.oddHeader.left.text = "营多天猫 · 经营总览"
    ws.page_setup.fitToHeight = 1
    ws.print_area = "A1:N40"
    return ws


def shop_formula(column_letter, aggregate="sum"):
    if aggregate == "sum":
        return f'IF(双店销售!$C$3="两家店",图表数据!{column_letter}2+图表数据!{column_letter}3,IF(双店销售!$C$3="旗舰店",图表数据!{column_letter}2,图表数据!{column_letter}3))'
    return ""


def write_stores(wb, refs):
    ws = wb.create_sheet("双店销售", 1)
    setup_dashboard(ws, TEAL_BG, "0F766E")
    for row in range(1, 47):
        for col in range(1, 17):
            ws.cell(row, col).fill = fill(TEAL_BG)
            ws.cell(row, col).font = font(10, color=TEAL_MUTED)
    ws.merge_cells("B2:M2")
    title = ws.cell(2, 2, "双店销售 · 7月27日–8月2日")
    title.font = font(22, True, "F3FFFB")
    ws.cell(3, 2, "店铺").font = font(11, True, TEAL_MUTED)
    ws.merge_cells("C3:D3")
    picker = ws.cell(3, 3, "两家店")
    picker.font = font(12, True, "134E4A")
    picker.fill = fill("ECFDF8")
    picker.alignment = align("center")
    ws.cell(3, 4).fill = fill("ECFDF8")
    picker.comment = Comment("选择两家店、旗舰店或淘宝店。指标卡和三张图都会跟着变。", "仪表盘")
    validation = DataValidation(type="list", formula1='"两家店,旗舰店,淘宝店"', allow_blank=False)
    validation.add(ws["C3"])
    ws.add_data_validation(validation)
    ws.merge_cells("E3:M3")
    ws.cell(3, 5, "口径与字段复核一致：已发货待签收、已签收、已发货退款成功、已签收退款成功。").font = font(9, False, TEAL_MUTED)

    money = '#,##0'
    cards = [
        (2, "商家实收", f'={shop_formula("D")}', money, [
            ("订单数", f'={shop_formula("B")}', "#,##0"),
            ("买家实付", f'={shop_formula("E")}', money),
            ("实收占比", '=IF($C$3="两家店",图表数据!D2/图表数据!D4,IF($C$3="旗舰店",图表数据!D2/图表数据!D4,图表数据!D3/图表数据!D4))', "0.0%"),
        ]),
        (5, "成交件数", f'={shop_formula("C")}', "#,##0", [
            ("旗舰件数", "=图表数据!C2", "#,##0"),
            ("淘宝件数", "=图表数据!C3", "#,##0"),
            ("件均实收", f'=IF({shop_formula("C")}=0,NA(),{shop_formula("D")}/{shop_formula("C")})', "0.00"),
        ]),
        (8, "总成本", f'={shop_formula("G")}', money, [
            ("占去税收入", f'=IF({shop_formula("F")}=0,NA(),{shop_formula("G")}/{shop_formula("F")})', "0.0%"),
            ("去税收入", f'={shop_formula("F")}', money),
            ("旗舰成本", "=图表数据!G2", money),
        ]),
        (11, "推广前利润", f'={shop_formula("H")}', money, [
            ("净利率", f'=IF({shop_formula("F")}=0,NA(),{shop_formula("H")}/{shop_formula("F")})', "0.0%"),
            ("旗舰利润", "=图表数据!H2", money),
            ("淘宝利润", "=图表数据!H3", money),
        ]),
    ]
    for col, title_text, hero, hero_format, lines in cards:
        write_card(
            ws, col, title_text, hero, hero_format, lines,
            TEAL_CARD, "115E59", muted=TEAL_MUTED, hero_color="F3FFFB",
        )

    ws.cell(12, 2, "每日商家实收").font = font(11, True, "F3FFFB")
    ws.cell(12, 8, "省份商家实收").font = font(11, True, "F3FFFB")
    source = wb["图表数据"]
    add_lines(
        ws, "B13", "每日商家实收",
        Reference(source, min_col=10, min_row=refs["daily_first"], max_row=refs["daily_last"]),
        Reference(source, min_col=11, min_row=refs["daily_header"], max_col=12, max_row=refs["daily_last"]),
    )
    add_clustered(
        ws, "H13", "省份",
        Reference(source, min_col=1, min_row=refs["prov_first"], max_row=refs["prov_last"]),
        Reference(source, min_col=4, min_row=refs["prov_header"], max_col=5, max_row=refs["prov_last"]),
        ["5EEAD4", "F9A8D4"],
        horizontal=True,
        width=13.2,
        height=8.0,
        largest_first=True,
    )
    ws.cell(30, 2, "系列商家实收").font = font(11, True, "F3FFFB")
    add_clustered(
        ws, "B31", "系列",
        Reference(source, min_col=1, min_row=refs["series_first"], max_row=refs["series_last"]),
        Reference(source, min_col=4, min_row=refs["series_header"], max_col=5, max_row=refs["series_last"]),
        ["5EEAD4", "F9A8D4"],
        horizontal=True,
        width=22,
        height=7.6,
        largest_first=True,
    )
    ws.merge_cells("B46:M46")
    ws.row_dimensions[46].height = 32
    foot = ws.cell(
        46, 2,
        "两家店视图的实收占比是旗舰店占两家店的比例。选中单店时，占比是该店占两家店的比例，图中只保留该店。"
        "商家实收 115,659，总成本 50,371，与字段复核对齐。未支付取消和未发货退款不在这个口径里。",
    )
    foot.font = font(9, False, TEAL_MUTED)
    foot.alignment = align("left", "center", True)
    ws.print_area = "A1:N46"
    ws.oddHeader.left.text = "营多天猫 · 双店销售"
    return ws


def style_header(ws, headers, widths):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    ws.row_dimensions[1].height = 22
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    for col, name in enumerate(headers, 1):
        cell = ws.cell(1, col, name)
        cell.fill = fill(HEADER)
        cell.font = font(10, True, WHITE)
        cell.alignment = align("center")
        ws.column_dimensions[get_column_letter(col)].width = widths.get(name, 14)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def write_skus(wb, weeks):
    ws = wb.create_sheet("商品周损益")
    ws.sheet_properties.tabColor = "7C3AED"
    headers = [
        "周次", "商品ID", "定位", "分组", "商品名称",
        "成交金额", "去税金额", "成交件数", "件单价",
        "万相台花费", "万相台成交", "万相台ROI",
        "品销宝花费", "品销宝成交", "品销宝ROI",
        "推广总花费", "推广总成交", "推广ROI",
        "日销成交", "其他费用",
        "推广前利润", "推广前净利率", "推广后利润", "推广后净利率",
    ]
    widths = {"商品ID": 18, "定位": 16, "分组": 14, "商品名称": 42, "周次": 14}
    style_header(ws, headers, widths)
    percent = {"推广前净利率", "推广后净利率"}
    roi = {"万相台ROI", "品销宝ROI", "推广ROI"}
    money = {"成交金额", "去税金额", "件单价", "万相台花费", "万相台成交", "品销宝花费", "品销宝成交", "推广总花费", "推广总成交", "日销成交", "其他费用", "推广前利润", "推广后利润"}
    row_index = 2
    for week in weeks:
        for sku in week["skus"]:
            values = [
                week["label"], sku["id"], sku["role"], sku["group"], sku["name"],
                sku["gmv"], sku["net"], sku["qty"], sku["asp"],
                sku["wxtSpend"], sku["wxtSales"], sku["wxtRoi"],
                sku["pxbSpend"], sku["pxbSales"], sku["pxbRoi"],
                sku["adSpend"], sku["adSales"], sku["adRoi"],
                sku["organic"], sku["other"],
                sku["profitBefore"], sku["marginBefore"], sku["profitAfter"], sku["marginAfter"],
            ]
            band = fill(ZEBRA) if row_index % 2 == 0 else fill("FFFFFF")
            for col, (name, value) in enumerate(zip(headers, values), 1):
                cell = ws.cell(row_index, col, value)
                cell.font = font(9, color=INK)
                cell.fill = band
                cell.alignment = align("left" if col <= 5 else "right")
                if name in percent:
                    cell.number_format = '0.0"%"'
                elif name in roi and value is not None:
                    cell.number_format = "0.00"
                elif name in money:
                    cell.number_format = '#,##0.00'
                elif name == "成交件数":
                    cell.number_format = "#,##0"
                elif name == "商品ID":
                    cell.number_format = "@"
            row_index += 1
    last = row_index - 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"
    red = Font(name="微软雅黑", size=9, color="BE123C")
    for letter in ("V", "X"):
        ws.conditional_formatting.add(f"{letter}2:{letter}{last}", CellIsRule(operator="lessThan", formula=["0"], font=red))
    ws.auto_filter.ref = f"A1:X{last}"
    ws.oddHeader.left.text = "营多天猫 · 商品周损益"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return last - 1


def write_orders(wb, orders):
    ws = wb.create_sheet("订单明细")
    ws.sheet_properties.tabColor = "0F766E"
    headers = ["店铺", "支付日", "商品ID", "商品标题", "品类", "系列", "省", "订单状态", "计入损益", "件数", "买家实付", "商家实收", "总成本", "推广前利润"]
    widths = {"商品ID": 18, "商品标题": 42, "订单状态": 18, "支付日": 14, "计入损益": 12, "系列": 18, "品类": 10}
    style_header(ws, headers, widths)
    for row_index, record in enumerate(orders["rows"], 2):
        shop, sku_id, title, category, series, province, status, qty, pay, recv, cost, profit, day = record
        included = "是" if status in VALID_STATUS else "否"
        values = [shop, day, sku_id, title, category, series, province, status, included, qty, pay, recv, cost, profit]
        band = fill("F3FFFB") if row_index % 2 == 0 else fill("FFFFFF")
        for col, value in enumerate(values, 1):
            cell = ws.cell(row_index, col, value)
            cell.font = font(9, color=INK)
            cell.fill = band
            cell.alignment = align("left" if col <= 9 else "right")
            if col == 2 and value:
                cell.number_format = "YYYY-MM-DD"
            elif col == 3:
                cell.number_format = "@"
            elif col == 10:
                cell.number_format = "#,##0"
            elif col >= 11:
                cell.number_format = '#,##0.00'
    last = len(orders["rows"]) + 1
    ws.auto_filter.ref = f"A1:N{last}"
    red = Font(name="微软雅黑", size=9, color="BE123C")
    ws.conditional_formatting.add(f"N2:N{last}", CellIsRule(operator="lessThan", formula=["0"], font=red))
    ws.oddHeader.left.text = "营多天猫 · 订单明细"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.auto_filter.ref = f"A1:N{last}"
    return len(orders["rows"])


def main():
    weeks = [load_week(*item) for item in WEEKS]
    orders = load_orders()
    rows = week_rows(weeks)
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    write_week_sheet(wb, rows)
    refs = write_chart_data(wb, weeks, orders, week_last=1 + len(weeks))
    write_overview(wb, refs)
    write_stores(wb, refs)
    sku_count = write_skus(wb, weeks)
    order_count = write_orders(wb, orders)
    wb.defined_names.add(DefinedName(name="周次列表", attr_text=f"周汇总!$A$2:$A${1 + len(weeks)}"))
    wb.active = wb["经营总览"]
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.properties.title = "营多捞面天猫经营仪表盘"
    wb.properties.creator = "营多天猫仪表盘"
    order_recv = round(sum(record[9] for record in orders["rows"] if record[6] in VALID_STATUS), 2)
    order_cost = round(sum(record[10] for record in orders["rows"] if record[6] in VALID_STATUS), 2)
    if abs(order_recv - 115659.23) > 0.05 or abs(order_cost - 50371.1) > 0.05:
        raise SystemExit(f"order totals drifted recv={order_recv} cost={order_cost}")
    if abs(rows[-1]["成交金额"] - 107470) > 1:
        raise SystemExit(f"unexpected last week gmv {rows[-1]['成交金额']}")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT}")
    print(f"weeks {len(weeks)} skus {sku_count} orders {order_count} recv {order_recv} cost {order_cost}")


if __name__ == "__main__":
    main()

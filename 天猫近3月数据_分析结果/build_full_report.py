#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import pandas as pd
import numpy as np
import re
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, Reference
from openpyxl.utils import get_column_letter
import shutil

src = Path("/workspace/新建文件夹 (2)")
out_dir = Path("/workspace/天猫近3月数据_分析结果")
out_dir.mkdir(exist_ok=True)


def clean_col(c):
    return re.sub(r"\s+", "", str(c).replace("\n", "").strip())


def parse_money(x):
    if pd.isna(x):
        return np.nan
    s = str(x).strip()
    if s in ("", "-", "—", "nan", "None", "#DIV/0!", "#N/A"):
        return np.nan
    s = s.replace("￥", "").replace(",", "").replace("%", "").replace(" ", "")
    if s.startswith("(") and s.endswith(")"):
        s = "-" + s[1:-1]
    try:
        return float(s)
    except Exception:
        return np.nan


def week_key(w):
    start = str(w).split("-")[0]
    m, d = start.split(".")
    return (int(m), int(d))


rows = []
files = sorted(set(list(src.glob("天猫*.csv")) + list(src.glob("天猫，*.csv"))))
for f in files:
    if any(k in f.name for k in ["净利润", "趋势", "归因", "环比", "领导", "拆解", "更新"]):
        continue
    try:
        df = pd.read_csv(f, encoding="utf-8-sig")
    except Exception:
        df = pd.read_csv(f, encoding="gbk")
    df.columns = [clean_col(c) for c in df.columns]
    if "商品ID" not in df.columns:
        continue
    for _, r in df.iterrows():
        date_in = str(r.get("日期", "")).strip()
        if date_in in ("总计", "nan", "") or date_in.lower() == "nan":
            continue
        pid = r.get("商品ID")
        if pd.isna(pid):
            continue
        if "7.6-7.12" in f.name:
            week = "7.6-7.12"
        elif "7.20-7.26" in f.name:
            week = "7.20-7.26"
        elif "7.27-8.2" in f.name and "周报" not in f.name:
            week = date_in
        else:
            week = date_in
        rows.append(
            {
                "week": week,
                "id": str(int(float(pid))),
                "定位": str(r.get("定位", "")).strip(),
                "商品名称": str(r.get("商品名称", "")).strip(),
                "成交金额": parse_money(r.get("成交金额")),
                "推广总花费": parse_money(r.get("推广总花费")),
                "净利润_推广前": parse_money(r.get("净利润(推广前)")),
                "净利润_推广后": parse_money(r.get("净利润(推广后)")),
            }
        )

prod = pd.DataFrame(rows).drop_duplicates(["week", "id"], keep="last")
weekly = prod[prod["week"] != "7.20-8.2"].copy()

overall = (
    weekly.groupby("week", as_index=False)
    .agg(
        成交金额=("成交金额", "sum"),
        净利润_推广前=("净利润_推广前", "sum"),
        推广总花费=("推广总花费", "sum"),
        净利润_推广后=("净利润_推广后", "sum"),
        SKU数=("id", "nunique"),
    )
    .sort_values("week", key=lambda s: s.map(week_key))
)
overall["推广后环比"] = overall["净利润_推广后"].diff()
overall["推广前环比"] = overall["净利润_推广前"].diff()
overall["推广费环比"] = overall["推广总花费"].diff()


def decomp(a_w, b_w):
    common = set(weekly.loc[weekly["week"] == a_w, "id"]) & set(
        weekly.loc[weekly["week"] == b_w, "id"]
    )
    out = []
    for wid in common:
        ra = weekly[(weekly["id"] == wid) & (weekly["week"] == a_w)].iloc[0]
        rb = weekly[(weekly["id"] == wid) & (weekly["week"] == b_w)].iloc[0]
        d_after = rb["净利润_推广后"] - ra["净利润_推广后"]
        d_before = rb["净利润_推广前"] - ra["净利润_推广前"]
        d_ad = rb["推广总花费"] - ra["推广总花费"]
        effect_before, effect_ad = d_before, -d_ad
        residual = d_after - (d_before - d_ad)
        if d_after >= 0:
            cause = "本期改善/持平"
        else:
            hurt_before = max(0, -effect_before)
            hurt_ad = max(0, -effect_ad)
            if hurt_ad == 0 and hurt_before == 0:
                cause = "其他费用/口径差"
            elif hurt_ad >= hurt_before * 1.2:
                cause = "主要是推广费上升"
            elif hurt_before >= hurt_ad * 1.2:
                cause = "主要是推广前利润下降"
            else:
                cause = "推广前利润下降 + 推广费上升（双因）"
        out.append(
            {
                "商品ID": wid,
                "定位": rb["定位"],
                "商品名称": rb["商品名称"],
                "成交变化": rb["成交金额"] - ra["成交金额"],
                "推广前净利变化": d_before,
                "推广花费变化": d_ad,
                "推广后净利变化": d_after,
                "推广前影响": effect_before,
                "推广费影响": effect_ad,
                "其他影响": residual,
                "主因判断": cause,
            }
        )
    return pd.DataFrame(out).sort_values("推广后净利变化")


d1 = decomp("7.6-7.12", "7.13-7.19")
d2 = decomp("7.13-7.19", "7.20-7.26")
d3 = decomp("7.6-7.12", "7.20-7.26")

best = overall.loc[overall["week"] == "7.6-7.12"].iloc[0]
w713 = overall.loc[overall["week"] == "7.13-7.19"].iloc[0]
w720 = overall.loc[overall["week"] == "7.20-7.26"].iloc[0]
m1 = d1[d1["商品ID"] == "668566625747"].iloc[0]
m2 = d2.iloc[0]
m2_main = d2[d2["商品ID"] == "668566625747"]
w1 = d1[d1["推广后净利变化"] < 0].head(8)
w2 = d2[d2["推广后净利变化"] < 0].head(8)
w3 = d3[d3["推广后净利变化"] < 0].head(10)


def cause_summary(df):
    w = df[df["推广后净利变化"] < 0]
    parts = []
    for cause, g in w.groupby("主因判断"):
        parts.append(f"{cause}{len(g)}个/{g['推广后净利变化'].sum():.0f}")
    return "；".join(parts)


NAVY, BLUE, ORANGE, GREEN, RED, WHITE, LIGHT_BLUE, SOFT_RED = (
    "1F4E79",
    "2E75B6",
    "C65911",
    "548235",
    "C00000",
    "FFFFFF",
    "D6EAF8",
    "FCE4D6",
)
thin = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)


def hdr(ws, row, sc, ec, color=NAVY):
    for c in range(sc, ec + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = PatternFill("solid", fgColor=color)
        cell.font = Font(name="微软雅黑", bold=True, color=WHITE, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin


def paint(ws, r1, r2, c1, c2, money=None):
    money = money or set()
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name="微软雅黑", size=10)
            cell.border = thin
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
                horizontal="center" if c in money else "left",
            )
            if c in money and isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0"
                if cell.value < 0:
                    cell.font = Font(name="微软雅黑", size=10, color=RED)
                elif cell.value > 0:
                    cell.font = Font(name="微软雅黑", size=10, color=GREEN)


wb = Workbook()
ws = wb.active
ws.title = "01-完整报告"
ws.merge_cells("A1:H1")
ws["A1"] = "天猫近3个月利润下降原因分析报告（完整版）"
ws["A1"].font = Font(name="微软雅黑", bold=True, size=18, color=NAVY)
ws.merge_cells("A2:H2")
ws["A2"] = (
    "口径：净利润(推广后)｜覆盖周次：5.11–7.26（单周，已含7.20-7.26）｜单位：元｜"
    "拆解：推广后变化 ≈ 推广前变化 − 推广花费变化"
)
ws["A2"].font = Font(name="微软雅黑", size=9, color="666666")

lines = []
lines.append(
    f"总判断：整体长期亏损；7.6-7.12 改善至 {best['净利润_推广后']:.0f} 后连续回落，"
    f"7.13-7.19 为 {w713['净利润_推广后']:.0f}，7.20-7.26 进一步至 {w720['净利润_推广后']:.0f}。"
    f"相对最好周累计恶化 {w720['净利润_推广后']-best['净利润_推广后']:.0f} 元。"
)
lines.append(
    f"阶段一（7.6→7.13）：整体推广后 {best['净利润_推广后']:.0f}→{w713['净利润_推广后']:.0f}"
    f"（{w713['净利润_推广后']-best['净利润_推广后']:+.0f}）。主因是推广前利润下降，不是推广费上升。"
    f"第一责任ID：668566625747（袋面主链接），推广后 {m1['推广后净利变化']:+.0f}；"
    f"推广前 {m1['推广前净利变化']:+.0f}，推广费 {m1['推广花费变化']:+.0f}。"
)
extra = ""
if len(m2_main):
    mm = m2_main.iloc[0]
    extra = f"袋面主链接同期推广后 {mm['推广后净利变化']:+.0f}（已不再是最大拖累）。"
lines.append(
    f"阶段二（7.13→7.20）：整体 {w713['净利润_推广后']:.0f}→{w720['净利润_推广后']:.0f}"
    f"（{w720['净利润_推广后']-w713['净利润_推广后']:+.0f}）。本阶段推广费上升拖累更突出。"
    f"恶化最大ID：{m2['商品ID']}（{m2['定位']}），推广后 {m2['推广后净利变化']:+.0f}；"
    f"推广前 {m2['推广前净利变化']:+.0f}，推广费 {m2['推广花费变化']:+.0f} → {m2['主因判断']}。{extra}"
)
lines.append(f"阶段一归因结构（恶化ID）：{cause_summary(d1)}。第一波主要来自推广前利润走弱。")
lines.append(f"阶段二归因结构（恶化ID）：{cause_summary(d2)}。第二波中推广费上升成为重要推手。")
lines.append(
    "原味大大包：推广前常有利润，但推广后净贡献弱，不是两波回落的第一责任；"
    "部分大大包链接在阶段二也出现推广前下滑。"
)
lines.append(
    "长期亏损底座：袋面主/次链接、杯面主链接、巴东牛肉持续贡献大额负利润；"
    "近两周变化叠加其上，造成可见的利润下降。"
)
lines.append(
    "建议：①阶段一抓袋面主链接推广前利润/成交；②阶段二严控袋面次链接与30包推广费；"
    "③巴东牛肉高费比持续亏损建议减投；④大大包压推广侵蚀、保推广前利润。"
)

ws["A4"] = "一、利润下降原因（完整结论）"
ws["A4"].font = Font(name="微软雅黑", bold=True, size=13, color=NAVY)
ws["A5"] = "序号"
ws["B5"] = "说明"
hdr(ws, 5, 1, 2)
for i, line in enumerate(lines, 1):
    rr = 5 + i
    ws.cell(row=rr, column=1, value=i)
    ws.cell(row=rr, column=2, value=line)
    for c in range(1, 3):
        cell = ws.cell(row=rr, column=c)
        cell.font = Font(name="微软雅黑", size=10)
        cell.border = thin
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        if i % 2 == 0:
            cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    ws.row_dimensions[rr].height = 36

r = 15
ws.cell(row=r, column=1, value="二、整体周趋势").font = Font(
    name="微软雅黑", bold=True, size=13, color=NAVY
)
r += 1
headers = ["周次", "成交金额", "推广前净利润", "推广花费", "推广后净利润", "推广后环比", "说明"]
for j, h in enumerate(headers, 1):
    ws.cell(row=r, column=j, value=h)
hdr(ws, r, 1, 7)
r += 1
r0 = r
for _, row in overall.iterrows():
    if row["week"] == "7.6-7.12":
        note = "最好周"
    elif row["week"] == "7.13-7.19":
        note = "第一波回落"
    elif row["week"] == "7.20-7.26":
        note = "第二波回落/最新周"
    else:
        note = "单周"
    vals = [
        row["week"],
        row["成交金额"],
        row["净利润_推广前"],
        row["推广总花费"],
        row["净利润_推广后"],
        row["推广后环比"],
        note,
    ]
    for j, v in enumerate(vals, 1):
        ws.cell(row=r, column=j, value=None if pd.isna(v) else v)
    if row["week"] == "7.6-7.12":
        fill = PatternFill("solid", fgColor="C6EFCE")
    elif row["week"] in ("7.13-7.19", "7.20-7.26"):
        fill = PatternFill("solid", fgColor=SOFT_RED)
    else:
        fill = None
    if fill:
        for j in range(1, 8):
            ws.cell(row=r, column=j).fill = fill
    r += 1
paint(ws, r0, r - 1, 1, 7, money={2, 3, 4, 5, 6})

chart = LineChart()
chart.title = "推广后净利润周趋势"
chart.style = 10
chart.y_axis.title = "元"
data = Reference(ws, min_col=5, min_row=r0 - 1, max_row=r - 1)
cats = Reference(ws, min_col=1, min_row=r0, max_row=r - 1)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
chart.width = 16
chart.height = 8
ws.add_chart(chart, f"A{r + 1}")
r += 16

ws.cell(row=r, column=1, value="三、两阶段责任ID").font = Font(
    name="微软雅黑", bold=True, size=13, color=NAVY
)
r += 1
ws.cell(row=r, column=1, value="阶段一：7.6→7.13（主因：推广前利润下降）").font = Font(
    name="微软雅黑", bold=True, size=11, color=ORANGE
)
r += 1
h = ["排名", "商品ID", "定位", "推广前变化", "推广费变化", "推广后变化", "主因判断", "商品名称"]
for j, x in enumerate(h, 1):
    ws.cell(row=r, column=j, value=x)
hdr(ws, r, 1, 8, ORANGE)
r += 1
r1 = r
for i, row in w1.reset_index(drop=True).iterrows():
    vals = [
        i + 1,
        row["商品ID"],
        row["定位"],
        row["推广前净利变化"],
        row["推广花费变化"],
        row["推广后净利变化"],
        row["主因判断"],
        row["商品名称"],
    ]
    for j, v in enumerate(vals, 1):
        ws.cell(row=r, column=j, value=v)
    r += 1
paint(ws, r1, r - 1, 1, 8, money={4, 5, 6})

r += 1
ws.cell(row=r, column=1, value="阶段二：7.13→7.20（推广费上升拖累加重）").font = Font(
    name="微软雅黑", bold=True, size=11, color=ORANGE
)
r += 1
for j, x in enumerate(h, 1):
    ws.cell(row=r, column=j, value=x)
hdr(ws, r, 1, 8, ORANGE)
r += 1
r2 = r
for i, row in w2.reset_index(drop=True).iterrows():
    vals = [
        i + 1,
        row["商品ID"],
        row["定位"],
        row["推广前净利变化"],
        row["推广花费变化"],
        row["推广后净利变化"],
        row["主因判断"],
        row["商品名称"],
    ]
    for j, v in enumerate(vals, 1):
        ws.cell(row=r, column=j, value=v)
    r += 1
paint(ws, r2, r - 1, 1, 8, money={4, 5, 6})

r += 2
ws.cell(row=r, column=1, value="四、相对最好周到最新周累计恶化TOP").font = Font(
    name="微软雅黑", bold=True, size=13, color=NAVY
)
r += 1
for j, x in enumerate(
    ["排名", "商品ID", "定位", "推广前累计变化", "推广费累计变化", "推广后累计变化", "主因判断", "商品名称"],
    1,
):
    ws.cell(row=r, column=j, value=x)
hdr(ws, r, 1, 8, BLUE)
r += 1
r3 = r
for i, row in w3.reset_index(drop=True).iterrows():
    vals = [
        i + 1,
        row["商品ID"],
        row["定位"],
        row["推广前净利变化"],
        row["推广花费变化"],
        row["推广后净利变化"],
        row["主因判断"],
        row["商品名称"],
    ]
    for j, v in enumerate(vals, 1):
        ws.cell(row=r, column=j, value=v)
    r += 1
paint(ws, r3, r - 1, 1, 8, money={4, 5, 6})

ws.column_dimensions["A"].width = 10
ws.column_dimensions["B"].width = 18
ws.column_dimensions["C"].width = 14
ws.column_dimensions["D"].width = 14
ws.column_dimensions["E"].width = 14
ws.column_dimensions["F"].width = 14
ws.column_dimensions["G"].width = 28
ws.column_dimensions["H"].width = 52


def dump_decomp(ws, df, title):
    ws["A1"] = title
    ws["A1"].font = Font(name="微软雅黑", bold=True, size=14, color=NAVY)
    cols = [
        "商品ID",
        "定位",
        "商品名称",
        "成交变化",
        "推广前净利变化",
        "推广花费变化",
        "推广后净利变化",
        "推广前影响",
        "推广费影响",
        "其他影响",
        "主因判断",
    ]
    for j, h in enumerate(cols, 1):
        ws.cell(row=3, column=j, value=h)
    hdr(ws, 3, 1, len(cols))
    for i, row in df.reset_index(drop=True).iterrows():
        rr = 4 + i
        for j, h in enumerate(cols, 1):
            ws.cell(row=rr, column=j, value=None if pd.isna(row[h]) else row[h])
    paint(ws, 4, 3 + len(df), 1, len(cols), money=set(range(4, 11)))
    ws.auto_filter.ref = f"A3:K{3+len(df)}"
    ws.freeze_panes = "A4"
    for i, w in enumerate([16, 12, 50, 10, 12, 12, 12, 10, 10, 10, 26], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


ws2 = wb.create_sheet("02-整体周趋势")
ws2["A1"] = "整体周趋势明细"
ws2["A1"].font = Font(name="微软雅黑", bold=True, size=14, color=NAVY)
headers2 = [
    "周次",
    "成交金额",
    "推广前净利润",
    "推广前环比",
    "推广花费",
    "推广费环比",
    "推广后净利润",
    "推广后环比",
]
for j, h in enumerate(headers2, 1):
    ws2.cell(row=3, column=j, value=h)
hdr(ws2, 3, 1, 8)
for i, row in overall.reset_index(drop=True).iterrows():
    rr = 4 + i
    vals = [
        row["week"],
        row["成交金额"],
        row["净利润_推广前"],
        row["推广前环比"],
        row["推广总花费"],
        row["推广费环比"],
        row["净利润_推广后"],
        row["推广后环比"],
    ]
    for j, v in enumerate(vals, 1):
        ws2.cell(row=rr, column=j, value=None if pd.isna(v) else v)
paint(ws2, 4, 3 + len(overall), 1, 8, money=set(range(2, 9)))

dump_decomp(wb.create_sheet("03-阶段一拆解"), d1, "阶段一：7.6-7.12 → 7.13-7.19")
dump_decomp(wb.create_sheet("04-阶段二拆解"), d2, "阶段二：7.13-7.19 → 7.20-7.26")
dump_decomp(wb.create_sheet("05-累计拆解"), d3, "累计：7.6-7.12 → 7.20-7.26")

ws6 = wb.create_sheet("06-主链轨迹")
ws6["A1"] = "主亏ID多周轨迹"
ws6["A1"].font = Font(name="微软雅黑", bold=True, size=14, color=NAVY)
row = 3
for wid in [
    "668566625747",
    "649361890551",
    "649735171438",
    "648600900806",
    "1057479227624",
]:
    sub = weekly[weekly["id"] == wid].sort_values("week", key=lambda s: s.map(week_key))
    if sub.empty:
        continue
    ws6.cell(
        row=row,
        column=1,
        value=f"{wid}｜{sub['定位'].iloc[-1]}｜{sub['商品名称'].iloc[-1]}",
    ).font = Font(name="微软雅黑", bold=True, color=NAVY)
    row += 1
    for j, h in enumerate(["周次", "成交金额", "推广前净利润", "推广花费", "推广后净利润"], 1):
        ws6.cell(row=row, column=j, value=h)
    hdr(ws6, row, 1, 5)
    row += 1
    rs = row
    for _, x in sub.iterrows():
        for j, v in enumerate(
            [x["week"], x["成交金额"], x["净利润_推广前"], x["推广总花费"], x["净利润_推广后"]],
            1,
        ):
            ws6.cell(row=row, column=j, value=None if pd.isna(v) else v)
        if x["week"] in ("7.6-7.12", "7.13-7.19", "7.20-7.26"):
            for j in range(1, 6):
                ws6.cell(row=row, column=j).fill = PatternFill("solid", fgColor=SOFT_RED)
        row += 1
    paint(ws6, rs, row - 1, 1, 5, money={2, 3, 4, 5})
    row += 2

ws7 = wb.create_sheet("07-口径说明")
ws7["A1"] = "口径与阅读说明"
ws7["A1"].font = Font(name="微软雅黑", bold=True, size=14, color=NAVY)
notes = [
    ("指标", "默认以净利润(推广后)衡量利润；并用净利润(推广前)、推广总花费拆解下降原因。"),
    ("拆解公式", "推广后变化 ≈ 推广前变化 − 推广花费变化（存在少量其他费用/口径差）。"),
    (
        "主因规则",
        "恶化周内：推广前下降贡献显著大于推广费上升 → 推广前利润下降；反之 → 推广费上升；接近则双因。",
    ),
    ("数据周", "已纳入单周7.20-7.26；旧文件7.20-8.2为两周合计，本报告主趋势不采用。"),
    ("文件修正", "天猫ID周损益_7.6-7.12.csv 日期列曾误标为6.29-7.5，分析按文件名纠正。"),
    ("缺失", "仍缺部分历史周（如5.18-5.24等），长周期仅作参考。"),
]
ws7["A3"] = "项目"
ws7["B3"] = "说明"
hdr(ws7, 3, 1, 2)
for i, (a, b) in enumerate(notes, 1):
    ws7.cell(row=3 + i, column=1, value=a)
    ws7.cell(row=3 + i, column=2, value=b)
    for c in range(1, 3):
        cell = ws7.cell(row=3 + i, column=c)
        cell.font = Font(name="微软雅黑", size=10, bold=(c == 1))
        cell.border = thin
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws7.row_dimensions[3 + i].height = 36
ws7.column_dimensions["A"].width = 12
ws7.column_dimensions["B"].width = 90

out = out_dir / "天猫利润下降原因_完整报告.xlsx"
wb.save(out)
shutil.copy(out, src / out.name)

md = []
md.append("# 天猫近3个月利润下降原因分析报告")
md.append("")
md.append("**口径**：净利润（推广后）｜**期间**：5.11–7.26（单周）｜**单位**：元")
md.append("")
md.append("## 一、结论摘要")
md.append("")
md.append(
    f"1. 最好周为 **7.6-7.12（{best['净利润_推广后']:.0f}）**，之后连续两周回落："
)
md.append(
    f"   - 7.13-7.19：**{w713['净利润_推广后']:.0f}**（{w713['净利润_推广后']-best['净利润_推广后']:+.0f}）"
)
md.append(
    f"   - 7.20-7.26：**{w720['净利润_推广后']:.0f}**（较上周 {w720['净利润_推广后']-w713['净利润_推广后']:+.0f}；较最好周 {w720['净利润_推广后']-best['净利润_推广后']:+.0f}）"
)
md.append("2. **利润下降分两阶段、原因不同**：")
md.append(
    f"   - **阶段一（7.6→7.13）**：主要是 **推广前利润下降**。第一责任 ID：**668566625747（袋面主链接）**，"
    f"推广后 {m1['推广后净利变化']:+.0f}（推广前 {m1['推广前净利变化']:+.0f}，推广费 {m1['推广花费变化']:+.0f}）。"
)
md.append(
    f"   - **阶段二（7.13→7.20）**：更突出的是 **推广费上升**。恶化最大 ID：**{m2['商品ID']}（{m2['定位']}）**，"
    f"推广后 {m2['推广后净利变化']:+.0f}（推广前 {m2['推广前净利变化']:+.0f}，推广费 {m2['推广花费变化']:+.0f}）。"
)
md.append("3. 原味大大包：推广前有利润，但推广后净贡献弱，**不是两波回落的第一责任**。")
md.append("")
md.append("## 二、整体周趋势")
md.append("")
md.append("| 周次 | 成交金额 | 推广前净利润 | 推广花费 | 推广后净利润 | 环比 |")
md.append("|---|---:|---:|---:|---:|---:|")
for _, row in overall.iterrows():
    rb = row["推广后环比"]
    md.append(
        f"| {row['week']} | {row['成交金额']:.0f} | {row['净利润_推广前']:.0f} | "
        f"{row['推广总花费']:.0f} | {row['净利润_推广后']:.0f} | "
        f"{'' if pd.isna(rb) else f'{rb:.0f}'} |"
    )
md.append("")
md.append("## 三、阶段一责任榜（7.6→7.13，主因：推广前利润下降）")
md.append("")
md.append("| 排名 | 商品ID | 定位 | 推广前变化 | 推广费变化 | 推广后变化 | 主因 |")
md.append("|---:|---|---|---:|---:|---:|---|")
for i, row in w1.reset_index(drop=True).iterrows():
    md.append(
        f"| {i+1} | {row['商品ID']} | {row['定位']} | {row['推广前净利变化']:.0f} | "
        f"{row['推广花费变化']:.0f} | {row['推广后净利变化']:.0f} | {row['主因判断']} |"
    )
md.append("")
md.append("## 四、阶段二责任榜（7.13→7.20，推广费上升拖累加重）")
md.append("")
md.append("| 排名 | 商品ID | 定位 | 推广前变化 | 推广费变化 | 推广后变化 | 主因 |")
md.append("|---:|---|---|---:|---:|---:|---|")
for i, row in w2.reset_index(drop=True).iterrows():
    md.append(
        f"| {i+1} | {row['商品ID']} | {row['定位']} | {row['推广前净利变化']:.0f} | "
        f"{row['推广花费变化']:.0f} | {row['推广后净利变化']:.0f} | {row['主因判断']} |"
    )
md.append("")
md.append("## 五、建议")
md.append("")
md.append("1. 针对阶段一：复盘袋面主链接成交与推广前利润下滑。")
md.append("2. 针对阶段二：立刻压降袋面次链接、30包无效/低效推广费。")
md.append("3. 巴东牛肉持续高亏损，建议减投或停投验证。")
md.append("4. 大大包保留推广前利润，严控推广侵蚀。")
md.append("")
md.append("## 附件")
md.append("")
md.append("- Excel完整版：`天猫利润下降原因_完整报告.xlsx`")

md_path = out_dir / "天猫利润下降原因_完整报告.md"
md_path.write_text("\n".join(md), encoding="utf-8")
shutil.copy(md_path, src / md_path.name)
print("Wrote", out)
print("Wrote", md_path)
print("stage1", m1["商品ID"], m1["主因判断"], m1["推广后净利变化"])
print("stage2", m2["商品ID"], m2["主因判断"], m2["推广后净利变化"])
